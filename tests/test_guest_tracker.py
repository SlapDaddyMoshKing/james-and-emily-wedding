from io import BytesIO
import unittest
from unittest.mock import patch
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from botocore.exceptions import ClientError
from backend.guest_tracker import COLUMNS, SHEET, RECEIPTS, TRACKER_KEY, TrackerError, append_to_workbook, save_to_tracker

def fixture():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(list(COLUMNS))
    sheet.cell(2, 1, "Existing Guest")
    sheet.cell(2, 4, "Existing address")
    sheet.cell(3, 1).fill = PatternFill("solid", fgColor="FFEECC")
    sheet.cell(232, 1).number_format = "@"
    other = workbook.create_sheet("Notes")
    other.cell(1, 1, "=1+2")
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()

def guest(name="New Guest"):
    return {"submission_id": str(uuid4()), **{field: "" for field in COLUMNS.values()},
        "name_line_one": name, "address_line1": "123 Example", "city": "Example City",
        "region": "MA", "postal_code": "01234", "phone": "+15555550100"}

class TrackerTests(unittest.TestCase):
    def test_append_preserves_template_and_uses_next_actual_row(self):
        data = guest("=1+1")
        result = append_to_workbook(fixture(), data)
        book = load_workbook(BytesIO(result))
        sheet = book[SHEET]
        self.assertEqual(sheet['A2'].value, "Existing Guest")
        self.assertEqual(sheet['D2'].value, "Existing address")
        self.assertEqual(sheet['A3'].value, "=1+1")
        self.assertEqual(sheet['A3'].data_type, "s")
        self.assertEqual(sheet['H3'].value, "01234")
        self.assertEqual(sheet['I3'].value, "+15555550100")
        self.assertEqual(sheet['A3'].fill.fgColor.rgb, "00FFEECC")
        self.assertEqual(book['Notes']['A1'].value, "=1+2")
        self.assertEqual(book[RECEIPTS].sheet_state, "hidden")
        self.assertIsNone(append_to_workbook(result, data))
        second = load_workbook(BytesIO(append_to_workbook(result, guest())))
        self.assertEqual(second[SHEET]['A4'].value, "New Guest")
        with self.assertRaises(TrackerError):
            append_to_workbook(result, {**data, "city": "Changed"})

    def test_concurrent_update_reloads_workbook_without_losing_guest(self):
        class RacingS3:
            content = fixture()
            revision = 0
            writes = 0
            def get_object(self, **kwargs):
                return {"Body": BytesIO(self.content), "ETag": str(self.revision)}
            def put_object(self, **kwargs):
                self.writes += 1
                if self.writes == 1:
                    self.content = append_to_workbook(self.content, guest("Concurrent Guest"))
                    self.revision += 1
                if kwargs['IfMatch'] != str(self.revision):
                    raise ClientError({"Error": {"Code": "PreconditionFailed"}}, "PutObject")
                self.content = kwargs['Body']
                self.revision += 1
        s3 = RacingS3()
        data = guest()
        with patch('backend.guest_tracker.time.sleep'):
            save_to_tracker(s3, 'test', data)
        book = load_workbook(BytesIO(s3.content))
        self.assertEqual([book[SHEET].cell(row, 1).value for row in (2, 3, 4)], ['Existing Guest', 'Concurrent Guest', 'New Guest'])
        save_to_tracker(s3, 'test', data)
        self.assertEqual(s3.writes, 2)

    def test_changed_headers_fail_without_writing(self):
        book = load_workbook(BytesIO(fixture()))
        book[SHEET]['A1'] = 'Unexpected'
        stream = BytesIO()
        book.save(stream)
        with self.assertRaises(TrackerError):
            append_to_workbook(stream.getvalue(), guest())
