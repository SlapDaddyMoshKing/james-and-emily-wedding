"""Private, read-only owner viewer. Uses the local AWS profile; binds to loopback only."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import secrets
import sys
from threading import Thread
from urllib.request import urlopen
import webbrowser
from wsgiref.simple_server import make_server, WSGIRequestHandler

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.guest_list import DATA_DIR
from backend.guest_tracker import TRACKER_KEY, SHEET
from openpyxl import load_workbook

BUCKET = "wedding-site-guest-data-8f3d21"
HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Guest Tracker · Owner view</title>
<style>body{margin:0;background:#f6f8fa;color:#314d61;font:15px/1.6 system-ui,sans-serif}main{max-width:1500px;margin:auto;padding:32px}h1{font:38px Georgia,serif;margin:0}header{display:flex;gap:20px;align-items:center;justify-content:space-between;flex-wrap:wrap}button,a{font:inherit;padding:10px 16px;background:#314d61;color:white;border:0;border-radius:5px;cursor:pointer;text-decoration:none}nav{display:flex;gap:12px;flex-wrap:wrap}.table-wrap{overflow:auto;background:white;border:1px solid #d8e0e5;border-radius:6px}table{border-collapse:collapse;min-width:1100px;width:100%}td,th{padding:12px;text-align:left;border-bottom:1px solid #e4e9ed;white-space:pre-wrap;overflow-wrap:anywhere;max-width:280px}th{background:#e4eef5;position:sticky;top:0}#status{min-height:26px}.error{color:#a13928}.subtle{color:#5b6a75}button:focus-visible,a:focus-visible{outline:3px solid #b87f46;outline-offset:3px}#close{background:#e4eef5;color:#314d61}@media(max-width:600px){main{padding:20px}h1{font-size:30px}}</style></head>
<body><main><header><div><h1>Guest Tracker</h1><p class="subtle">Emily &amp; James · Private owner view</p></div>
<nav><button id="refresh">Refresh now</button><a href="download">Download latest Excel</a><button id="close">Close viewer</button></nav></header>
<p>This view reads the workbook in AWS and refreshes every 20 seconds. Downloaded Excel files are snapshots.</p>
<p id="status" role="status" aria-live="polite">Loading the latest workbook…</p>
<div class="table-wrap"><table><thead><tr id="headers"></tr></thead><tbody id="rows"></tbody></table></div></main>
<script>let pending=false;let timer;
async function refresh(){if(pending)return;pending=true;document.querySelector('#refresh').disabled=true;
try{const response=await fetch('data',{cache:'no-store',signal:AbortSignal.timeout(20000)});if(!response.ok)throw Error();const data=await response.json();
const head=document.querySelector('#headers');head.replaceChildren();for(const name of data.headers){const th=document.createElement('th');th.textContent=name;head.append(th);}
const rows=document.querySelector('#rows');rows.replaceChildren();for(const row of data.rows){const tr=document.createElement('tr');for(const value of row){const td=document.createElement('td');td.textContent=value;tr.append(td);}rows.append(tr);}
const status=document.querySelector('#status');status.className='';status.textContent=data.rows.length+' filled rows · Workbook updated '+new Date(data.updated_at).toLocaleString()+' · Checked '+new Date().toLocaleTimeString();
}catch{const status=document.querySelector('#status');status.className='error';status.textContent='Could not refresh from AWS. Check your connection and AWS profile. Any rows shown are from the last successful refresh.';}
finally{pending=false;document.querySelector('#refresh').disabled=false;}}
document.querySelector('#refresh').onclick=refresh;document.querySelector('#close').onclick=async()=>{clearInterval(timer);await fetch('close',{method:'POST'});document.querySelector('#status').textContent='Viewer stopped. You can close this tab.';document.querySelector('#refresh').disabled=true;};
refresh();timer=setInterval(refresh,20000);</script></body></html>"""

def create_owner_app(s3, token, shutdown=lambda: None):
    def app(environ, start_response):
        def respond(status, body, content_type="application/json", extra=()):
            if not isinstance(body, bytes):
                body = json.dumps(body).encode()
            start_response(status, [("Content-Type", content_type), ("Content-Length", str(len(body))),
                ("Cache-Control", "no-store"), ("Referrer-Policy", "no-referrer"),
                ("X-Content-Type-Options", "nosniff"), ("X-Frame-Options", "DENY"), *extra])
            return [body]
        host = environ.get("HTTP_HOST", "")
        if host.split(":")[0] != "127.0.0.1" or environ.get("REMOTE_ADDR") != "127.0.0.1":
            return respond("403 Forbidden", {"error": "Local owner access only."})
        origin = environ.get("HTTP_ORIGIN")
        if origin and origin != "http://" + host:
            return respond("403 Forbidden", {"error": "Local owner access only."})
        path = environ.get("PATH_INFO", "")
        prefix = "/" + token + "/"
        if not path.startswith(prefix):
            return respond("404 Not Found", {"error": "Not found."})
        route = path[len(prefix):]
        method = environ.get("REQUEST_METHOD", "GET")
        if route == "close" and method == "POST":
            Thread(target=shutdown, daemon=True).start()
            return respond("200 OK", {"closed": True})
        if method != "GET":
            return respond("405 Method Not Allowed", {"error": "Use GET."})
        if route == "":
            return respond("200 OK", HTML.encode(), "text/html; charset=utf-8")
        if route == "health":
            return respond("200 OK", {"app": "wedding-owner-viewer"})
        if route not in {"data", "download"}:
            return respond("404 Not Found", {"error": "Not found."})
        try:
            response = s3.get_object(Bucket=BUCKET, Key=TRACKER_KEY)
            content = response["Body"].read()
            if route == "download":
                return respond("200 OK", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    [("Content-Disposition", 'attachment; filename="Guest Tracker.xlsx"')])
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook[SHEET]
            headers = [str(cell.value or "").strip() for cell in sheet[1]][:9]
            rows = [[str(value) if value is not None else "" for value in row]
                    for row in sheet.iter_rows(min_row=2, max_col=9, values_only=True)
                    if any(value is not None and value != "" for value in row)]
            workbook.close()
            return respond("200 OK", {"headers": headers, "rows": rows, "updated_at": response["LastModified"].isoformat()})
        except Exception:
            return respond("503 Service Unavailable", {"error": "Could not read the workbook. Check your AWS profile and connection."})
    return app

class QuietHandler(WSGIRequestHandler):
    def log_message(self, *args):
        pass  # No access logs containing the private viewer URL.

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args()
    state = DATA_DIR / "owner-viewer.json"
    try:
        existing = json.loads(state.read_text())
        port = existing["port"]
        token = existing["token"]
        if not isinstance(port, int) or not 1 <= port <= 65535 or not isinstance(token, str) or not token.isalnum():
            raise ValueError()
        url = f"http://127.0.0.1:{port}/{token}/"
        with urlopen(url + "health", timeout=2) as response:
            if json.load(response) != {"app": "wedding-owner-viewer"}:
                raise ValueError()
        if not args.no_browser:
            webbrowser.open(url)
        return
    except (OSError, ValueError, KeyError, TypeError):
        pass
    import boto3
    s3 = boto3.Session(profile_name="wedding-site", region_name="us-east-2").client("s3")
    token = secrets.token_hex(32)
    server = make_server("127.0.0.1", 0, lambda e, s: [], handler_class=QuietHandler)
    server.set_app(create_owner_app(s3, token, server.shutdown))
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    state.write_text(json.dumps({"port": server.server_port, "token": token}), encoding="utf-8")
    if not args.no_browser:
        webbrowser.open(f"http://127.0.0.1:{server.server_port}/{token}/")
    try:
        server.serve_forever()
    finally:
        server.server_close()
        state.unlink(missing_ok=True)

if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        (DATA_DIR / "owner-viewer-error.txt").write_text(str(error), encoding="utf-8")
        if sys.platform == "win32":
            import ctypes
            ctypes.windll.user32.MessageBoxW(0, "Could not start the Guest Tracker. Check owner-viewer-error.txt in your WeddingSiteData folder.", "Guest Tracker", 0)
        raise
