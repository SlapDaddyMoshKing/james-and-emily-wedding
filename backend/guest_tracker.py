"""Append one guest to the existing private Excel tracker, with atomic S3 retries."""
from copy import copy
from hashlib import sha256
from io import BytesIO
import json
import time
from zipfile import BadZipFile

from openpyxl import load_workbook
from botocore.exceptions import ClientError

TRACKER_KEY = "welcome/Guest Tracker.xlsx"
SHEET = "GUEST ADDRESSING"
RECEIPTS = "_WebsiteSubmissions"
COLUMNS = {
    "Name Line One": "name_line_one", "Name Line Two": "name_line_two",
    "Inner Envelope": "inner_envelope", "Address Line One": "address_line1",
    "Address Line Two": "address_line2", "City": "city", "State": "region",
    "Zip Code": "postal_code", "Phone Number": "phone",
}

class TrackerError(ValueError):
    pass

def append_to_workbook(content, data):
    """Return updated bytes, or None if this exact submission is already present."""
    try:
        workbook = load_workbook(BytesIO(content))
    except (BadZipFile, KeyError, ValueError) as error:
        raise TrackerError("The guest tracker could not be read.") from error
    if SHEET not in workbook:
        raise TrackerError("Guest addressing sheet is missing.")
    sheet = workbook[SHEET]
    headers = [str(sheet.cell(1, column).value or "").strip() for column in range(1, 10)]
    if headers != list(COLUMNS):
        raise TrackerError("Guest tracker columns have changed.")
    digest = sha256(json.dumps({field: data[field] for field in COLUMNS.values()}, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    if RECEIPTS in workbook:
        receipts = workbook[RECEIPTS]
        if tuple(cell.value for cell in receipts[1]) != ("Submission ID", "Payload SHA256", "Guest row"):
            raise TrackerError("Submission references have changed.")
        for identifier, fingerprint, _ in receipts.iter_rows(min_row=2, max_col=3, values_only=True):
            if identifier == data["submission_id"]:
                if fingerprint != digest:
                    raise TrackerError("Submission reference is already in use.")
                return None
    else:
        receipts = workbook.create_sheet(RECEIPTS)
        receipts.append(["Submission ID", "Payload SHA256", "Guest row"])
        receipts.sheet_state = "hidden"
    # The provided template has hundreds of styled empty rows. Start after the
    # last row containing actual content, without overwriting existing guests.
    last = max((row[0].row for row in sheet.iter_rows() if any(cell.value is not None and cell.value != "" for cell in row)), default=1)
    target = last + 1
    for column, field in enumerate(COLUMNS.values(), 1):
        cell = sheet.cell(target, column)
        if not cell.has_style and target > 2:
            cell._style = copy(sheet.cell(target - 1, column)._style)
        cell.value = data[field]
        cell.data_type = "s"  # Literal text, including leading =, +, or @.
        if field in {"postal_code", "phone"}:
            cell.number_format = "@"
    receipts.append([data["submission_id"], digest, target])
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f"A1:I{max(target, sheet.max_row)}"
    for table in sheet.tables.values():
        from openpyxl.utils.cell import range_boundaries
        left, top, right, bottom = range_boundaries(table.ref)
        if left == 1 and top == 1 and right == 9 and target > bottom:
            table.ref = f"A1:I{target}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()

def save_to_tracker(s3, bucket, data, key=TRACKER_KEY):
    for attempt in range(4):
        response = s3.get_object(Bucket=bucket, Key=key)
        updated = append_to_workbook(response["Body"].read(), data)
        if updated is None:
            return
        try:
            s3.put_object(Bucket=bucket, Key=key, Body=updated,
                ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ServerSideEncryption="AES256", IfMatch=response["ETag"],
                Metadata=response.get("Metadata", {}))
            return
        except ClientError as error:
            if error.response.get("Error", {}).get("Code") not in {"PreconditionFailed", "ConditionalRequestConflict"}:
                raise
            if attempt < 3:
                time.sleep(0.1 * (attempt + 1))
    raise TrackerError("Guest tracker is busy. Please retry the same submission.")
